# -*- coding: utf-8 -*-
"""
Martech'in kendi fiyat listesini (MARTECH EKRAN 20.09.2026.xlsx) TeknikERP
stok yukleme bicimine cevirir.

Kaynak dosya: her sayfa bir marka grubu; her sayfada birden fazla blok
(baslik satiri = marka adi + FIYAT + KOD). Sutun dizilisi sayfaya gore
degisir; LAYOUT sozlugu bunu sayfa adiyla acikca tanimlar. Bloklarin
altinda fiyatsiz "uyumluluk" listeleri var (Cin tedarikci kataloglari);
bunlar urun degildir, atlanir.

Kullanim:
    python tools/martech-excel-donustur.py "<kaynak.xlsx>" "<hedef.xlsx>"

Cikti sayfalari:
    Stoklar  - TeknikERP'nin okudugu sayfa (ilk sayfa olmali)
    Notlar   - varsayimlar, atlanan/isaretli satirlar
    Kaynak   - Stoklar satiri -> kaynak hucre (izlenebilirlik)
    LCD Uyum - musterinin uyumluluk notlari, temizlenmis (bilgi)
"""
import os
import re
import sys
from collections import Counter, OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KAYNAK = sys.argv[1]
HEDEF = sys.argv[2]

KATEGORI = "EKRAN & LCD"

# TeknikERP disa aktarma basliklari (backend/src/utils/excelExchange.ts STOK_BASLIKLARI)
BASLIKLAR = ["Id", "StokKodu", "StokAdi", "Kategori", "Marka", "Model", "Gorunum", "Kalite",
             "Renk", "Aciklama", "Rmb", "AlisFiyati", "Satis1", "Satis2", "AlisAdedi",
             "SatisAdedi", "Bakiye", "Uyumlu", "GelenAdet"]

# Sayfa -> sutun dizilisi
LAYOUT = {
    "APPLE": dict(model="C", uretici="D", kalite="E", coz="F", fiyat="G", kod="H", ek=["I"]),
    "SAMSUNG QUİKTEL": dict(model="B", gorunum="C", kalite="D", fiyat="E", kod="F", ek=["G"]),
    "HUAWEİ-ONE PLS-LENOVO": dict(model="B", gorunum="C", kalite="D", fiyat="E", kod="F", ek=["G"]),
    "XİAOMİ REDMİ-REEDER": dict(marka="B", model="C", gorunum="D", kalite="E", fiyat="F", kod="G", ek=["H"]),
    "OPPO NUBİA": dict(model="B", gorunum="C", kalite="D", fiyat="E", kod="F", ek=["G"]),
    "TEKNO-İNFİNİX-VİVO-LG-MEİZU MTR": dict(model="B", gorunum="C", kalite="D", fiyat="E", kod="F", ek=["G"]),
    "CSPR GM MİPO OMX TCL TRDN ": dict(model="B", gorunum="C", fiyat="D", kod="E", ek=["F"]),
}

# Blok basligi (kaynaktaki yazi, buyuk harf) -> (Marka, model on eki, not)
BLOK_MARKA = {
    "APPLE & IPHONE": ("APPLE", "İPHONE ", ""),
    "SAMSUNG": ("SAMSUNG", "", ""),
    "SAMSUNG TABLET FULL": ("SAMSUNG", "TAB ", "Tablet"),
    "OUKİTEL SERVİS": ("OUKITEL", "", ""),
    "HUAWEİ": ("HUAWEI", "", ""),
    "ONE PLUS": ("ONEPLUS", "", ""),
    "LENOVO": ("LENOVO", "", ""),
    "XİAOMİ & REDMİ": ("XIAOMI", "", ""),
    "REEDER LCD": ("REEDER", "", ""),
    "OPPO": ("OPPO", "", ""),
    "NUBİA": ("NUBIA", "", ""),
    "İNFİNİX LCD": ("INFINIX", "", ""),
    "TECNO LCD": ("TECNO", "", ""),
    "VİVO SERVİS H.K & OLED": ("VIVO", "", ""),
    "VİVO": ("VIVO", "", ""),
    "LG LCD": ("LG", "", ""),
    "MEİZU & MOTO": ("MEIZU", "", ""),
    "CASPER": ("CASPER", "", ""),
    "GENERAL MOBİLE": ("GENERAL MOBILE", "", ""),
    "OMİX LCD": ("OMIX", "", ""),
    "TCL": ("TCL", "", ""),
    "HİKİNG LCD": ("HIKING", "", ""),
    "TRİDENT LCD": ("TRIDENT", "", ""),
    "MİPO": ("MIPO", "", ""),
    "WİKO": ("WIKO", "", ""),
    # uyumluluk listesi basliklari - altindaki satirlar fiyatsiz, zaten atlanir
    "TEKNO İNFİNİX": ("INFINIX", "", ""),
}

# Xiaomi sayfasindaki alt marka sutunu
XIAOMI_ALT = {"Mİ": ("XIAOMI", "MI "), "XİAOMİ": ("XIAOMI", ""), "XİOAMİ": ("XIAOMI", ""),
              "REDMİ": ("XIAOMI", "REDMI "), "REDM": ("XIAOMI", "REDMI "), "POCO": ("XIAOMI", "POCO "),
              "REEDER": ("REEDER", "")}

# Model basina yapismis 3 harfli etiketler (raf/tedarikci kodu oldugu saniliyor - teyit bekliyor)
ETIKETLER = ["KRX", "RRX", "RTX", "ZRT", "PPX", "PTX", "FPX", "OTK", "CCX", "RPX", "NMK", "TTX", "UYX",
             "JCK", "YYX", "YYL", "LMX", "PCX", "HGX", "NKX", "PSX", "RJX", "TJX", "OPX", "GRX", "HMX",
             "DDX", "HHX"]

# Model metninden cikarilip nota tasinan ifadeler
MODEL_NOTLARI = [
    ("TAM BOYUT", "Tam Boyut"), ("BÜYÜK BOYUT", "Büyük Boyut"), ("( SMALL )", "Small"), ("(BİG)", "Big"),
    ("LCD FULL", "LCD Full"), ("LCD SERVİS", ""), ("LCD  SERVİS", ""), ("FULL  SERVİS", ""),
    ("TÜM VERSİYONLAR", "Tüm Versiyonlar"), ("( ÇİFT FİLM )", "Çift Film"),
]
# Model metninde kalite soyleyen kelimeler: kalite bos ise oradan dolar
MODEL_KALITE = [
    ("REVİZE ORJ.", "Revizyon Orijinal"), ("( ORG. )", "Orijinal"), ("ORJ.", "Orijinal"), ("ORG.", "Orijinal"),
    (" ORG", "Orijinal"), ("( TFT )", "TFT"), ("( OLED )", "Oled"), (" OLED", "Oled"), (" TFT", "TFT"),
    ("SERVİS", "HK Servis"),
]

GORUNUM_MAP = OrderedDict([
    ("ÇITASIZ", "Çıtasız"), ("ÇITALI", "Çıtalı"), ("KASALI", "Kasalı"),
])

KALITE_MAP = {
    "HK SERVİS": "HK Servis", "HK": "HK Servis", "HK SERVİS HD+": "HK Servis HD+", "SERVİS": "HK Servis",
    "OLED": "Oled", "OLED TAM BOYUT": "Oled", "HARD OLED": "Hard Oled", "OLED HARD": "Hard Oled",
    "SOFT OLED": "Soft Oled", "OLED SOFT": "Soft Oled",
    "ORJİNAL": "Orijinal", "ORJ": "Orijinal", "ORG": "Orijinal", "ORGİNAL": "Orijinal",
    "REVİZE": "Revizyon Orijinal", "FOG ÇİN ORJ": "Fog Çin Orijinal",
    "ÇIKMA": "Çıkma Orijinal", "ÇIKMA ORJ": "Çıkma Orijinal", "İKİNCİ EL DİAGNSOTİK": "İkinci El Diagnostik",
    "AAA": "AAA", "AA": "AA", "L.W": "",
    "TFT & İNCEL": "TFT & İncell", "TFT": "TFT", "TFT HD+": "TFT HD+",
    "İNCELL HD+": "İncell HD+", "İNCELL FHD+ COF": "İncell FHD+ COF", "İNCELL FHD+": "İncell FHD+",
    "İNCELL FHD": "İncell FHD", "FHD+ COF": "İncell FHD+ COF", "HD": "HD", "FHD": "FHD",
}


def temiz(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def buyuk(s):
    """Turkce buyuk harf: i -> İ (Python .upper() 'I' verir)."""
    return s.replace("i", "İ").upper()


def kalite_cevir(raw, notlar):
    raw = temiz(raw)
    if not raw:
        return ""
    diag = False
    r = raw
    if "DİAGNOSTİK" in r:
        diag = True
        r = temiz(r.replace("DİAGNOSTİK", ""))
    if r == "OLED TAM BOYUT":
        notlar.append("Tam Boyut")
    if r in KALITE_MAP:
        out = KALITE_MAP[r]
    else:
        out = r
        notlar.append(f"Kalite tanınmadı: {raw}")
    if diag and out:
        out += " Diagnostik"
    return out


def gorunum_cevir(raw):
    """Gorunum sutunu bazen kalite de tasiyor ('HK ÇITASIZ', 'ÇITASIZ OLED'). (gorunum, kalan) doner."""
    raw = temiz(raw)
    if not raw:
        return "", ""
    g = ""
    kalan = raw
    if "LCD FULL" in raw:
        return "", "LCD FULL"
    for k, v in GORUNUM_MAP.items():
        if k in raw:
            g = v
            kalan = temiz(raw.replace(k, ""))
            break
    kalan = temiz(kalan.replace("TAM BOYUT", ""))
    return g, kalan


def sayi(v):
    """'12,5' -> 12.5; sayi degilse None"""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parcalar(model_str):
    """'A52 & A72/A92' -> ['A52', 'A72', 'A92']"""
    p = re.split(r"\s*(?:/|&|,| - )\s*", model_str)
    return [x.strip(" .") for x in p if x.strip(" .")]


def fmt(x):
    if x is None:
        return ""
    return int(x) if float(x).is_integer() else x


# --------------------------------------------------------------------------
wb = load_workbook(KAYNAK, data_only=True)
urunler = []      # dict per product
kaynaklar = []    # (sheet, row, model metni)
atlanan = Counter()
uyari = []        # (kaynak, mesaj)

for ws in wb.worksheets:
    if ws.title not in LAYOUT:
        continue
    lay = LAYOUT[ws.title]
    marka, onek, blok_not = None, "", ""
    ek_cols = lay.get("ek", [])

    for r in range(1, ws.max_row + 1):
        def v(key):
            c = lay.get(key)
            return ws[f"{c}{r}"].value if c else None

        fiyat_raw = v("fiyat")
        # Blok basligi: FIYAT yazan satir; marka adi B veya C sutununda
        if temiz(fiyat_raw).upper() == "FİYAT":
            ad = ""
            for c in ("B", "C"):
                ad = temiz(ws[f"{c}{r}"].value)
                if ad and ad.upper() != "FİYAT":
                    break
            key = buyuk(ad)
            if key in BLOK_MARKA:
                marka, onek, blok_not = BLOK_MARKA[key]
            else:
                uyari.append((f"{ws.title}!{r}", f"Blok başlığı tanınmadı: {ad}"))
            continue

        model_raw = v("model")
        gor_raw = v("gorunum")
        kal_raw = v("kalite")
        kod_raw = v("kod")
        uret_raw = v("uretici")
        coz_raw = v("coz")

        # Yalniz marka adi yazan satir (NUBİA, VİVO ...) - blok degistirir
        if model_raw and not any(x not in (None, "") for x in (gor_raw, kal_raw, fiyat_raw, kod_raw, uret_raw)):
            key = buyuk(temiz(model_raw))
            if key in BLOK_MARKA:
                marka, onek, blok_not = BLOK_MARKA[key]
                continue
            atlanan["fiyatsız/özelliksiz satır (uyumluluk listesi vb.)"] += 1
            continue

        g_norm, g_kalan = gorunum_cevir(gor_raw)
        urun_mu = (fiyat_raw not in (None, "") or kod_raw not in (None, "") or g_norm
                   or (kal_raw not in (None, "") and "kalite" in lay) or uret_raw not in (None, ""))
        if not urun_mu or not temiz(model_raw):
            if temiz(model_raw) or temiz(gor_raw):
                atlanan["fiyatsız/özelliksiz satır (uyumluluk listesi vb.)"] += 1
            continue
        if marka is None:
            uyari.append((f"{ws.title}!{r}", "Marka bloğu bulunamadı, satır atlandı"))
            continue

        notlar = []
        acik = []
        if blok_not:
            notlar.append(blok_not)

        # ---- marka / model
        m = temiz(model_raw)
        row_marka, row_onek = marka, onek
        if "marka" in lay:
            alt = buyuk(temiz(v("marka")))
            if alt in XIAOMI_ALT:
                row_marka, row_onek = XIAOMI_ALT[alt]
            else:
                uyari.append((f"{ws.title}!{r}", f"Alt marka tanınmadı: {alt}"))

        etiket = ""
        mu = buyuk(m)
        for e in ETIKETLER:
            if mu.startswith(e) and len(m) > len(e):
                etiket = e
                m = m[len(e):].strip()
                break
        if etiket:
            acik.append(f"Etiket: {etiket}")

        for token, nt in MODEL_NOTLARI:
            if token in m:
                m = temiz(m.replace(token, " "))
                if nt:
                    notlar.append(nt)
        kalite_modelden = ""
        for token, kq in MODEL_KALITE:
            mu = buyuk(m)
            if token in mu:
                idx = mu.find(token)
                m = temiz(m[:idx] + " " + m[idx + len(token):])
                kalite_modelden = kalite_modelden or kq
        m = temiz(m.replace("LCD", " ")).strip(" .-")
        m = m.replace("POCCO", "POCO")
        if row_marka == "APPLE":
            m = re.sub(r"(\d)(PRO|PLUS|MAX|MİNİ)", r"\1 \2", m)
        if not m:
            uyari.append((f"{ws.title}!{r}", f"Model boş kaldı: {model_raw!r}"))
            continue

        up = buyuk(m)
        # marka sozcugu modelin basinda ise ayir
        for bw, bm in (("HONOR ", "HONOR"), ("REALME ", "REALME"), ("RELAME ", "REALME"), ("ONE PLUS ", "ONEPLUS"),
                       ("TCL ", "TCL"), ("REEDER ", "REEDER"), ("İNF ", "INFINIX"), ("SAMSUNG ", "SAMSUNG")):
            if up.startswith(bw):
                row_marka = bm
                m = m[len(bw):].strip()
                break

        parts = parcalar(m)
        model_ilk = row_onek + parts[0]
        model_tam = row_onek + m
        uyumlu = []
        for p in parts[1:]:
            uyumlu.append((row_onek if row_marka == "APPLE" else "") + p)

        # ---- gorunum / kalite
        gorunum = g_norm
        kal_notlar = []
        kalite = kalite_cevir(kal_raw, kal_notlar) if "kalite" in lay else ""
        if g_kalan == "LCD FULL" or buyuk(temiz(kal_raw)) == "LCD FULL":
            notlar.append("LCD Full")
            if buyuk(temiz(kal_raw)) == "LCD FULL":
                kalite = ""
            g_kalan = ""
        if g_kalan == "REVİZE":
            kalite = "Revizyon Orijinal"
            g_kalan = ""
        if not kalite and g_kalan:
            kalite = kalite_cevir(g_kalan, kal_notlar)
        elif g_kalan:
            kal_notlar.append(f"Görünüm sütununda ek: {g_kalan}")
        if not kalite and kalite_modelden:
            kalite = kalite_modelden
        for kn in kal_notlar:
            if kn.startswith("Kalite tanınmadı") or kn.startswith("Görünüm"):
                uyari.append((f"{ws.title}!{r}", kn))
            else:
                notlar.append(kn)

        uretici = buyuk(temiz(uret_raw)) if uret_raw else ""
        if uretici and uretici != "ORJ":
            acik.append(f"Üretici: {uretici}")
        else:
            uretici = ""
        coz = buyuk(temiz(coz_raw))
        if coz:
            acik.append(f"Çözünürlük: {coz}")

        # ---- fiyat / kod
        fiyat = sayi(fiyat_raw)
        if fiyat_raw not in (None, "") and fiyat is None:
            uyari.append((f"{ws.title}!{r}", f"Fiyat okunamadı: {fiyat_raw!r}"))
        if fiyat is None:
            uyari.append((f"{ws.title}!{r}", "Fiyat boş - 0 yazıldı"))
            fiyat = 0
        kod = temiz(kod_raw)
        if kod:
            acik.append(f"Kod: {kod}")

        # ---- ek sutunlar (uyumlu modeller / notlar)
        for c in ek_cols:
            ev = ws[f"{c}{r}"].value
            if ev in (None, ""):
                continue
            evs = temiz(ev)
            if sayi(ev) is not None or evs in ("??", ".", ".."):
                uyari.append((f"{ws.title}!{c}{r}", f"Ek hücre anlaşılmadı: {evs}"))
                continue
            uyumlu += parcalar(evs)

        # ---- stok adi
        ad_parca = [("İPHONE" if row_marka == "APPLE" else row_marka),
                    m if row_marka == "APPLE" else model_tam,
                    "LCD", buyuk(gorunum) if gorunum else "", buyuk(kalite) if kalite else "", uretici]
        ek_ad = []
        if coz:
            ek_ad.append(coz)
        for nt in notlar:
            ek_ad.append(buyuk(nt))
        stok_adi = " ".join(x for x in ad_parca if x)
        if ek_ad:
            stok_adi += " (" + ", ".join(ek_ad) + ")"

        aciklama = " · ".join(acik + [n for n in notlar if n != "Tablet"])

        urunler.append(dict(
            StokAdi=stok_adi, Kategori=KATEGORI, Marka=row_marka, Model=model_ilk,
            Gorunum=gorunum, Kalite=kalite, Renk="", Aciklama=aciklama, Rmb="", AlisFiyati="",
            Satis1=fmt(fiyat), Satis2=fmt(fiyat), AlisAdedi="", SatisAdedi="", Bakiye=0,
            Uyumlu=", ".join(OrderedDict.fromkeys(u for u in uyumlu if u)), GelenAdet="",
        ))
        kaynaklar.append((ws.title, r, temiz(model_raw)))

# ---- mukerrer ad
sayac = Counter(u["StokAdi"] for u in urunler)
gorulen = Counter()
mukerrer = []
for u, k in zip(urunler, kaynaklar):
    ad = u["StokAdi"]
    if sayac[ad] > 1:
        gorulen[ad] += 1
        if gorulen[ad] > 1:
            u["StokAdi"] = f"{ad} ({gorulen[ad]})"
            mukerrer.append((f"{k[0]}!{k[1]}", f"Aynı ad: {ad}"))

# ---- LCD UYUM sayfasi (bilgi)
uyum = []
if "LCD UYUM" in wb.sheetnames:
    grup = ""
    for row in wb["LCD UYUM"].iter_rows(min_col=2, max_col=2, values_only=True):
        t = temiz(row[0])
        if not t or t in ("zt", "YK"):
            continue
        if re.match(r"^\d+\s*-", t):
            uyum.append((grup, re.sub(r"^\d+\s*-\s*", "", t)))
        elif "UYUMLU" in buyuk(t) or buyuk(t).startswith("TABLET"):
            grup = t
        else:
            uyum.append((grup, t))

# ---- yaz
out = Workbook()
ws1 = out.active
ws1.title = "Stoklar"
ws1.append(BASLIKLAR)
for u in urunler:
    ws1.append([u.get(h, "") for h in BASLIKLAR])
normal = Font(name="Arial")
bold = Font(name="Arial", bold=True)
for row in ws1.iter_rows():
    for c in row:
        c.font = normal
for c in ws1[1]:
    c.font = bold
    c.fill = PatternFill("solid", fgColor="DDEBF7")
ws1.freeze_panes = "C2"
ws1.auto_filter.ref = ws1.dimensions
genis = {"StokAdi": 62, "Kategori": 14, "Marka": 16, "Model": 28, "Gorunum": 10, "Kalite": 24,
         "Aciklama": 44, "Uyumlu": 40}
for i, h in enumerate(BASLIKLAR, 1):
    ws1.column_dimensions[get_column_letter(i)].width = genis.get(h, 10)

ws2 = out.create_sheet("Notlar")
markalar = Counter(u["Marka"] for u in urunler)
fiyatsiz = sum(1 for u in urunler if u["Satis1"] in (0, "", None))
notlar_genel = [
    ("Kaynak", os.path.basename(KAYNAK)),
    ("Ürün satırı", len(urunler)),
    ("Atlanan (fiyatsız uyumluluk/katalog satırları)", sum(atlanan.values())),
    ("Fiyatı boş olan ürün (Satis1 = 0)", fiyatsiz),
    ("Aynı ada düşüp (2), (3) eki alan", len(mukerrer)),
    ("", ""),
    ("VARSAYIMLAR", ""),
    ("Kategori", f"Hepsi '{KATEGORI}' - dosya yalnızca ekran listesi"),
    ("Fiyat", "FİYAT sütunu = Satış 1 = Satış 2 (dolar varsayıldı; teyit edilecek)"),
    ("KOD sütunu", "Anlamı teyit edilmedi; olduğu gibi Açıklama'ya 'Kod: …' olarak yazıldı. Alış/RMB doldurulmadı."),
    ("3 harfli ön ekler (KRX, PPX, JCK…)", "Model adından ayrıldı, Açıklama'ya 'Etiket: …' yazıldı. Anlamı teyit edilecek."),
    ("Apple üretici kodu (ZY, GX, DD, LW)", "Stok adına ve Açıklama'ya yazıldı ('Üretici: ZY')."),
    ("Xiaomi / Redmi / Poco", "Marka XIAOMI, model 'REDMI NOTE 9' / 'POCO X3' / 'MI 9T' (Shenzhen düzeniyle aynı)."),
    ("Honor / Realme / OnePlus", "Kendi markası olarak ayrıldı (HUAWEI / OPPO bloğundan)."),
    ("Çoklu model ('A52 & A72/A92')", "Model = ilk ad; kalanı Uyumlu sütununda; stok adı tam metni taşır."),
    ("'TAM BOYUT', '( SMALL )' vb.", "Modelden çıkarılıp stok adına parantezle ve Açıklama'ya not olarak eklendi."),
    ("Stok / maliyet", "Dosyada yok: Bakiye 0, AlisFiyati boş. StokKodu boş → sistem EKR00001… üretir."),
    ("Renk", "Dosyada yok, boş bırakıldı."),
    ("", ""),
    ("MARKA DAĞILIMI", ""),
] + [(k, v) for k, v in markalar.most_common()] + [("", ""), ("UYARILAR (kaynak hücre → not)", "")] + uyari + mukerrer
for a, b in notlar_genel:
    ws2.append([a, b])
ws2.column_dimensions["A"].width = 44
ws2.column_dimensions["B"].width = 110
for row in ws2.iter_rows():
    for c in row:
        c.font = normal
        c.alignment = Alignment(wrap_text=True, vertical="top")

ws3 = out.create_sheet("Kaynak")
ws3.append(["Stoklar satırı", "StokAdi", "Kaynak sayfa", "Kaynak satır", "Kaynak model metni"])
for i, (u, (s, r, mr)) in enumerate(zip(urunler, kaynaklar), start=2):
    ws3.append([i, u["StokAdi"], s, r, mr])
for col, w in zip("ABCDE", (14, 62, 34, 12, 44)):
    ws3.column_dimensions[col].width = w
for row in ws3.iter_rows():
    for c in row:
        c.font = normal

ws4 = out.create_sheet("LCD Uyum")
ws4.append(["Grup", "Birbirinin yerine kullanılan modeller (müşterinin notu)"])
for g, t in uyum:
    ws4.append([g, t])
ws4.column_dimensions["A"].width = 40
ws4.column_dimensions["B"].width = 120
for row in ws4.iter_rows():
    for c in row:
        c.font = normal
        c.alignment = Alignment(wrap_text=True, vertical="top")
for wsx in (ws2, ws3, ws4):
    for c in wsx[1]:
        c.font = bold

out.save(HEDEF)
print(f"urun={len(urunler)} atlanan={dict(atlanan)} fiyatsiz={fiyatsiz} mukerrer={len(mukerrer)} uyari={len(uyari)}")
print("markalar:", markalar.most_common())
